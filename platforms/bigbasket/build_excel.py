import json, datetime, statistics, os
from collections import defaultdict, OrderedDict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

d = json.load(open('result.json'))
rows = d['allRows']
per = d['perPin']
summary = d['summary']

# platform name derived from the folder this runs in (blinkit, zepto, ...)
PLATFORM = os.path.basename(os.getcwd()).replace('-', ' ').title()

JIVO_GREEN = "008B3A"
HDR = PatternFill("solid", fgColor=JIVO_GREEN)
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, color=JIVO_GREEN, size=18)
SUB_FONT = Font(italic=True, color="555555", size=10)
RED = PatternFill("solid", fgColor="F4CCCC")
GREEN = PatternFill("solid", fgColor="D9EAD3")
YEL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")

wb = Workbook()

# pretty SKU label from canonical
def label(canon):
    parts = canon.rsplit('-', 1)
    name = parts[0].replace('-', ' ').title()
    pack = parts[1].upper().replace('ML', ' ml').replace('L', ' L') if len(parts) > 1 else ''
    return f"{name} {pack}".strip()

skus = sorted(set(r['canonical'] for r in rows))
cities_with = sorted(set(r['city'] for r in rows))
all_cities = OrderedDict()
for p in per:
    all_cities.setdefault(p['city'], 0)
    all_cities[p['city']] += len(p['rows'])
cities_without = [c for c, n in all_cities.items() if n == 0]

def style_header(ws, row=1, ncols=None):
    ncols = ncols or ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR; cell.font = HDR_FONT; cell.alignment = CEN; cell.border = BORDER

def autosize(ws, maxw=42):
    for col in ws.columns:
        L = get_column_letter(col[0].column)
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[L].width = min(maxw, max(10, w + 2))

# ---------- Sheet 1: Executive Summary ----------
ws = wb.active; ws.title = "Summary"
ws["A1"] = f"Jivo x {PLATFORM} - Live Pricing Intelligence"; ws["A1"].font = TITLE_FONT
ws.merge_cells("A1:G1")
ws["A2"] = f"Captured {summary['captured_at'][:16].replace('T',' ')} IST  -  {summary['pincodes_with_jivo']}/{summary['pincodes_total']} pincodes carry Jivo  -  {summary['unique_skus']} unique SKUs  -  {summary['total_rows']} datapoints  -  scrape {summary['wall_s']}s"
ws["A2"].font = SUB_FONT; ws.merge_cells("A2:G2")

# KPI cards
kpis = [("Unique SKUs", summary['unique_skus']), ("Pincodes w/ Jivo", f"{summary['pincodes_with_jivo']}/{summary['pincodes_total']}"),
        ("Datapoints", summary['total_rows']), ("Cities w/ ZERO Jivo", len(cities_without))]
r = 4
for i, (k, v) in enumerate(kpis):
    c = 1 + i * 2
    ws.cell(row=r, column=c, value=k).font = Font(bold=True, size=10, color="555555")
    cell = ws.cell(row=r + 1, column=c, value=v); cell.font = Font(bold=True, size=20, color=JIVO_GREEN)

# cheapest per SKU
ws.cell(row=7, column=1, value="Cheapest pincode per SKU (in-stock, sale price)").font = Font(bold=True, size=12)
hdr = ["SKU", "City", "Pincode", "Store", "Sale Rs", "MRP Rs", "Disc %"]
for j, h in enumerate(hdr, 1):
    ws.cell(row=8, column=j, value=h)
style_header(ws, 8, len(hdr))
rr = 9
for s in skus:
    cand = [x for x in rows if x['canonical'] == s and x['in_stock'] == 1]
    if not cand: cand = [x for x in rows if x['canonical'] == s]
    if not cand: continue
    b = min(cand, key=lambda x: x['sale'])
    for j, v in enumerate([label(s), b['city'], b['pincode'], b['store_name'], b['sale'], b['mrp'], b['discount_pct']], 1):
        cell = ws.cell(row=rr, column=j, value=v); cell.border = BORDER
        if j >= 5: cell.alignment = CEN
    rr += 1
# distribution gaps
rr += 1
ws.cell(row=rr, column=1, value=f"WHITESPACE - Cities with ZERO Jivo on {PLATFORM}:").font = Font(bold=True, size=11, color="CC0000")
rr += 1
ws.cell(row=rr, column=1, value=", ".join(cities_without) if cities_without else "None - full coverage").font = Font(size=11)
ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=7)
autosize(ws)

# ---------- Sheet 2: Master Data ----------
ws = wb.create_sheet("Master Data")
cols = ["City", "Pincode", "Locality", "Store", "SKU", "Pack", "Vol (ml)", "Sale Rs", "MRP Rs", "Disc %", "Rs/L", "ETA min", "In stock"]
ws.append(cols)
for x in sorted(rows, key=lambda r: (r['city'], r['pincode'], r['canonical'])):
    ws.append([x['city'], x['pincode'], x['locality'], x['store_name'], x['sku_raw'], x['pack'], x['vol_ml'],
               x['sale'], x['mrp'], x['discount_pct'], x['per_litre'], x['eta_min'], "Yes" if x['in_stock'] else "No"])
style_header(ws)
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.border = BORDER
        if cell.column in (8, 9): cell.number_format = '"Rs"#,##0'
        if cell.column == 10: cell.number_format = '0.0"%"'
    sc = row[7].value
    if row[12].value == "No": row[12].fill = RED
    if isinstance(row[9].value, (int, float)) and row[9].value and row[9].value >= 40: row[9].fill = GREEN
autosize(ws)

# ---------- Matrix builder ----------
def matrix(sheet_name, valfn, fmt=None, scale=False, scale_rev=False):
    ws = wb.create_sheet(sheet_name)
    cols = ["SKU"] + cities_with
    ws.append(cols)
    for s in skus:
        rowvals = [label(s)]
        for c in cities_with:
            cand = [x for x in rows if x['canonical'] == s and x['city'] == c]
            rowvals.append(valfn(cand) if cand else None)
        ws.append(rowvals)
    style_header(ws)
    ws.freeze_panes = "B2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            if cell.column > 1:
                cell.alignment = CEN
                if fmt and isinstance(cell.value, (int, float)): cell.number_format = fmt
                if cell.value is None: cell.value = "-"
    if scale:
        last = get_column_letter(ws.max_column)
        lo, hi = ("F4CCCC", "D9EAD3") if scale_rev else ("D9EAD3", "F4CCCC")
        ws.conditional_formatting.add(f"B2:{last}{ws.max_row}",
            ColorScaleRule(start_type="min", start_color=lo, end_type="max", end_color=hi))
    autosize(ws, maxw=14)
    return ws

# Sheet 3: Pricing Matrix (modal sale per city) - green cheap -> red expensive
def modal_price(cands):
    prices = [x['sale'] for x in cands if x['sale'] is not None]
    if not prices: return None
    cnt = Counter(prices); top = max(cnt.values())
    return round(min(p for p, n in cnt.items() if n == top))
matrix("Pricing Matrix", modal_price, '"Rs"#,##0', scale=True)
# Sheet 4: Stock Status (% in stock)
def stock_cell(c):
    pct = round(100 * sum(x['in_stock'] for x in c) / len(c))
    return pct
wsS = matrix("Stock Status", stock_cell, '0"%"')
for row in wsS.iter_rows(min_row=2):
    for cell in row:
        if cell.column > 1 and isinstance(cell.value, (int, float)):
            cell.fill = GREEN if cell.value == 100 else (RED if cell.value == 0 else YEL)
# Sheet 5: Discount Analysis (modal disc per city) - higher = greener
def modal_disc(cands):
    vals = [x['discount_pct'] for x in cands if x['discount_pct'] is not None]
    if not vals: return None
    cnt = Counter(vals); top = max(cnt.values())
    return round(min(v for v, n in cnt.items() if n == top), 1)
matrix("Discount Analysis", modal_disc, '0.0"%"', scale=True, scale_rev=True)

# ---------- Sheet 6: Coverage / Gaps ----------
ws = wb.create_sheet("Coverage & Gaps")
ws.append(["City", "Pincode", "Locality", "Store assigned", "Jivo SKUs found"])
for p in per:
    ws.append([p['city'], p['pincode'], p['locality'], p['store_name'], len(p['rows'])])
style_header(ws); ws.freeze_panes = "A2"
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.border = BORDER
    if row[4].value == 0: row[4].fill = RED
    else: row[4].fill = GREEN
autosize(ws)

fname = f"Jivo-{PLATFORM.replace(' ', '')}-Live-Report-{datetime.date.today()}.xlsx"
wb.save(fname)
print("SAVED:", fname)
print("Sheets:", wb.sheetnames)
